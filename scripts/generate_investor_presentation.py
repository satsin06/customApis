#!/usr/bin/env python3
"""Generate a consolidated investor presentation from three source decks."""

from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable, Sequence

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT_DIR = Path(__file__).resolve().parents[1]
SOURCE_PAGES_DIR = ROOT_DIR / "presentation_assets" / "source_pages"
CONTACT_SHEETS_DIR = ROOT_DIR / "presentation_assets" / "contact_sheets"
OUTPUT_PPTX = ROOT_DIR / "Factory_Analytics_Investor_Presentation.pptx"

APP_PREFIX = "Factory_Analytics_App_Architecture"
BACKEND_PREFIX = "Factory_Analytics_Backend_Architecture_Flow"
WEB_PREFIX = "Factory_Analytics_Web_App_Architecture"

SOURCE_PDFS = {
    APP_PREFIX: ROOT_DIR / "Factory Analytics App Architecture.pdf",
    BACKEND_PREFIX: ROOT_DIR / "Factory Analytics Backend Architecture Flow.pdf",
    WEB_PREFIX: ROOT_DIR / "Factory Analytics Web App Architecture.pdf",
}

BG_COLOR = RGBColor(6, 15, 37)
PANEL_COLOR = RGBColor(15, 28, 55)
TEXT_COLOR = RGBColor(235, 240, 250)
MUTED_TEXT_COLOR = RGBColor(170, 185, 215)
ACCENT_COLOR = RGBColor(0, 184, 212)


def list_source_images(prefix: str) -> list[Path]:
    images = sorted(SOURCE_PAGES_DIR.glob(f"{prefix}_p*.png"))
    if len(images) != 13:
        _render_source_pdf(prefix)
        images = sorted(SOURCE_PAGES_DIR.glob(f"{prefix}_p*.png"))
        if len(images) != 13:
            raise RuntimeError(
                f"Expected 13 slides for prefix '{prefix}', found {len(images)} in {SOURCE_PAGES_DIR}"
            )
    return images


def _render_source_pdf(prefix: str) -> None:
    pdf_path = SOURCE_PDFS.get(prefix)
    if not pdf_path:
        raise RuntimeError(f"No source PDF mapping found for '{prefix}'")
    if not pdf_path.exists():
        raise FileNotFoundError(f"Source PDF not found: {pdf_path}")

    SOURCE_PAGES_DIR.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(pdf_path))
    try:
        for idx in range(doc.page_count):
            page = doc.load_page(idx)
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            out_path = SOURCE_PAGES_DIR / f"{prefix}_p{idx + 1:02d}.png"
            pix.save(str(out_path))
    finally:
        doc.close()


def create_contact_sheet(images: Sequence[Path], output_path: Path) -> None:
    cols = 4
    thumb_w, thumb_h = 380, 214
    margin = 40
    title_h = 90
    rows = math.ceil(len(images) / cols)
    canvas_w = margin * 2 + cols * thumb_w + (cols - 1) * margin
    canvas_h = title_h + margin + rows * thumb_h + (rows - 1) * margin + margin

    sheet = Image.new("RGB", (canvas_w, canvas_h), color=(10, 20, 43))
    for idx, image_path in enumerate(images):
        row = idx // cols
        col = idx % cols
        x = margin + col * (thumb_w + margin)
        y = title_h + margin + row * (thumb_h + margin)
        with Image.open(image_path) as src:
            src = src.convert("RGB")
            src.thumbnail((thumb_w, thumb_h))
            paste_x = x + (thumb_w - src.width) // 2
            paste_y = y + (thumb_h - src.height) // 2
            sheet.paste(src, (paste_x, paste_y))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(output_path)


def _set_dark_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def _add_header(slide, title: str, subtitle: str | None = None) -> None:
    _set_dark_background(slide)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_COLOR
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(8.6), Inches(0.9))
    tf = tbox.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(38)
    run.font.bold = True
    run.font.color.rgb = TEXT_COLOR

    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.05), Inches(9.6), Inches(0.45))
        stf = sbox.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        srun = sp.add_run()
        srun.text = subtitle
        srun.font.size = Pt(18)
        srun.font.color.rgb = MUTED_TEXT_COLOR


def _add_bullets(
    slide,
    bullets: Iterable[str],
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int = 22,
) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL_COLOR
    box.line.color.rgb = RGBColor(40, 72, 118)
    box.line.width = Pt(1.25)

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(12)
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT_COLOR


def _add_picture_fit(slide, image_path: Path, left: float, top: float, width: float, height: float) -> None:
    with Image.open(image_path) as img:
        img_w, img_h = img.size
    box_w = Inches(width)
    box_h = Inches(height)
    img_ratio = img_w / img_h
    box_ratio = box_w / box_h

    if img_ratio > box_ratio:
        pic_w = box_w
        pic_h = int(box_w / img_ratio)
        offset_x = 0
        offset_y = int((box_h - pic_h) / 2)
    else:
        pic_h = box_h
        pic_w = int(box_h * img_ratio)
        offset_x = int((box_w - pic_w) / 2)
        offset_y = 0

    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), box_w, box_h)
    frame.fill.solid()
    frame.fill.fore_color.rgb = PANEL_COLOR
    frame.line.color.rgb = RGBColor(40, 72, 118)
    frame.line.width = Pt(1.0)

    slide.shapes.add_picture(
        str(image_path),
        Inches(left) + offset_x,
        Inches(top) + offset_y,
        width=pic_w,
        height=pic_h,
    )


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_dark_background(slide)

    tbox = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2))
    tf = tbox.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Factory Analytics"
    r1.font.size = Pt(62)
    r1.font.bold = True
    r1.font.color.rgb = TEXT_COLOR

    p2 = tf.add_paragraph()
    p2.space_before = Pt(8)
    r2 = p2.add_run()
    r2.text = "Investor Presentation"
    r2.font.size = Pt(32)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT_COLOR

    p3 = tf.add_paragraph()
    p3.space_before = Pt(24)
    r3 = p3.add_run()
    r3.text = "Unified operating intelligence for modern factories"
    r3.font.size = Pt(24)
    r3.font.color.rgb = MUTED_TEXT_COLOR

    foot = slide.shapes.add_textbox(Inches(0.9), Inches(6.8), Inches(11.5), Inches(0.4))
    ff = foot.text_frame
    ff.clear()
    fp = ff.paragraphs[0]
    fp.alignment = PP_ALIGN.RIGHT
    fr = fp.add_run()
    fr.text = "Consolidated from App, Backend, and Web Architecture decks"
    fr.font.size = Pt(14)
    fr.font.color.rgb = MUTED_TEXT_COLOR


def add_bullet_slide(prs: Presentation, title: str, subtitle: str, bullets: Sequence[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, title, subtitle)
    _add_bullets(slide, bullets, left=0.9, top=1.7, width=11.5, height=5.4, font_size=24)


def add_visual_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
    bullets: Sequence[str],
    image_path: Path,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, title, subtitle)
    _add_bullets(slide, bullets, left=0.7, top=1.65, width=5.6, height=5.55, font_size=19)
    _add_picture_fit(slide, image_path, left=6.55, top=1.65, width=6.1, height=5.55)


def add_two_column_slide(
    prs: Presentation,
    title: str,
    subtitle: str,
    left_title: str,
    left_bullets: Sequence[str],
    right_title: str,
    right_bullets: Sequence[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, title, subtitle)

    lt = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(5.5), Inches(0.4))
    ltf = lt.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    lrun = lp.add_run()
    lrun.text = left_title
    lrun.font.size = Pt(22)
    lrun.font.bold = True
    lrun.font.color.rgb = ACCENT_COLOR

    rt = slide.shapes.add_textbox(Inches(6.9), Inches(1.6), Inches(5.5), Inches(0.4))
    rtf = rt.text_frame
    rtf.clear()
    rp = rtf.paragraphs[0]
    rrun = rp.add_run()
    rrun.text = right_title
    rrun.font.size = Pt(22)
    rrun.font.bold = True
    rrun.font.color.rgb = ACCENT_COLOR

    _add_bullets(slide, left_bullets, left=0.8, top=2.0, width=5.7, height=4.9, font_size=19)
    _add_bullets(slide, right_bullets, left=6.8, top=2.0, width=5.7, height=4.9, font_size=19)


def add_contact_sheet_slide(prs: Presentation, title: str, subtitle: str, sheet_path: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, title, subtitle)
    _add_picture_fit(slide, sheet_path, left=0.8, top=1.6, width=11.8, height=5.7)


def build_deck() -> None:
    app_images = list_source_images(APP_PREFIX)
    backend_images = list_source_images(BACKEND_PREFIX)
    web_images = list_source_images(WEB_PREFIX)

    app_sheet = CONTACT_SHEETS_DIR / "app_architecture_contact_sheet.png"
    backend_sheet = CONTACT_SHEETS_DIR / "backend_architecture_contact_sheet.png"
    web_sheet = CONTACT_SHEETS_DIR / "web_architecture_contact_sheet.png"
    create_contact_sheet(app_images, app_sheet)
    create_contact_sheet(backend_images, backend_sheet)
    create_contact_sheet(web_images, web_sheet)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    add_bullet_slide(
        prs,
        title="Investment Thesis",
        subtitle="Why this product can scale in industrial environments",
        bullets=[
            "Factory Analytics digitizes production planning, execution, and quality into one operating system.",
            "Role-specific experiences drive daily adoption: mobile workflows for line teams and web analytics for leadership.",
            "A modular backend and standardized data flow create a scalable platform for multi-factory expansion.",
        ],
    )
    add_two_column_slide(
        prs,
        title="Problem and Strategic Opportunity",
        subtitle="Manufacturers need real-time visibility and operational discipline",
        left_title="Pain Points Today",
        left_bullets=[
            "Fragmented tools create blind spots between planning, line execution, and quality control.",
            "Manual reporting delays corrective actions, increasing rework and missed dispatch targets.",
            "Supervisors and operators work from different data sources, causing decision latency.",
        ],
        right_title="Opportunity With Factory Analytics",
        right_bullets=[
            "Single data model aligns floor operations with management decisions in near real time.",
            "Closed-loop quality workflows reduce defect leakage and improve accountability.",
            "Architecture supports phased rollout from one line to multi-site deployments.",
        ],
    )
    add_bullet_slide(
        prs,
        title="Platform Overview: Three Integrated Layers",
        subtitle="App architecture + backend flow + web architecture",
        bullets=[
            "Factory Floor App: role-based workflows for station setup, production updates, and quality capture.",
            "Backend Orchestration: secured request pipeline with validation, business services, and a generic data access layer.",
            "Web Control Tower: real-time dashboards, order segmentation, and lifecycle tracking for production leadership.",
            "Result: one operational truth from line input to executive KPI reporting.",
        ],
    )
    add_visual_slide(
        prs,
        title="Layer 1: Factory Floor App",
        subtitle="Source: Factory Analytics App Architecture",
        bullets=[
            "Designed for two critical personas: Line Incharge and Production Incharge.",
            "Captures production updates and quality evidence at the point of execution.",
            "Handles unstable connectivity with resilient sync and recovery patterns.",
            "Creates high-frequency operational data without increasing floor complexity.",
        ],
        image_path=app_images[4],
    )
    add_visual_slide(
        prs,
        title="Layer 2: Backend Architecture Flow",
        subtitle="Source: Factory Analytics Backend Architecture Flow",
        bullets=[
            "Every request moves through a strict route -> middleware -> service -> data pipeline.",
            "Security and validation guardrails enforce clean inputs before business logic execution.",
            "Layered responsibilities improve maintainability and accelerate feature iteration.",
            "Architecture is built for scale and predictable API behavior across clients.",
        ],
        image_path=backend_images[4],
    )
    add_visual_slide(
        prs,
        title="Layer 3: Web App Architecture",
        subtitle="Source: Factory Analytics Web App Architecture",
        bullets=[
            "Management dashboard turns live production data into actionable KPIs.",
            "Separates customer orders and self-production for clearer planning decisions.",
            "Supports real-time monitoring of efficiency, throughput, and fail rates.",
            "Serves as the command center for line-level and plant-level performance.",
        ],
        image_path=web_images[7],
    )
    add_visual_slide(
        prs,
        title="Industrial-Grade Reliability and Security",
        subtitle="Operational trust is core to product defensibility",
        bullets=[
            "Session continuity and token refresh prevent workflow disruption during shifts.",
            "Middleware-based auth and input validation protect critical production data.",
            "Global error handling and standardized responses simplify incident resolution.",
            "Resilience features are designed specifically for noisy, bandwidth-variable factory environments.",
        ],
        image_path=backend_images[5],
    )
    add_visual_slide(
        prs,
        title="Closed-Loop Operating Intelligence",
        subtitle="From planning to analysis in one connected lifecycle",
        bullets=[
            "Plan -> Resource -> Assembly -> Analyze becomes a measurable digital loop.",
            "Production events feed dashboards continuously, not only end-of-day summaries.",
            "Quality and throughput trends can be acted on before losses compound.",
            "This feedback loop is the foundation for margin expansion at scale.",
        ],
        image_path=web_images[9],
    )
    add_two_column_slide(
        prs,
        title="Business Model and Expansion",
        subtitle="Scalable commercialization path",
        left_title="Monetization",
        left_bullets=[
            "Subscription SaaS priced by factory site and active production lines.",
            "Implementation services for integration, onboarding, and role-based training.",
            "Premium analytics modules for advanced benchmarking and forecasting.",
        ],
        right_title="Expansion Path",
        right_bullets=[
            "Land with pilot line, then expand within plant and to sister facilities.",
            "Use role-based adoption data to drive upsell across departments.",
            "Build ecosystem integrations for ERP/MES interoperability over time.",
        ],
    )
    add_two_column_slide(
        prs,
        title="Execution Roadmap (Next 18 Months)",
        subtitle="Focused milestones for investor confidence",
        left_title="Product Milestones",
        left_bullets=[
            "H1: harden core workflows, role governance, and reliability telemetry.",
            "H2: launch benchmark dashboards and automated exception alerting.",
            "H3: add multi-site rollup analytics and enterprise admin controls.",
        ],
        right_title="Commercial Milestones",
        right_bullets=[
            "Pilot conversion playbook with quantified ROI by line and shift.",
            "Expand from pilot customers into multi-line annual contracts.",
            "Establish channel partnerships for faster manufacturing vertical penetration.",
        ],
    )
    add_bullet_slide(
        prs,
        title="Funding Use and Success Metrics",
        subtitle="Capital allocation linked to measurable outcomes",
        bullets=[
            "Product and engineering: reliability, analytics depth, and integration velocity.",
            "Customer success: deployment playbooks and measurable time-to-value reduction.",
            "Go-to-market: focused pipeline in manufacturing segments with repeatable motion.",
            "Core KPIs: deployment cycle time, weekly active users, defect trend reduction, and revenue retention.",
        ],
    )
    add_bullet_slide(
        prs,
        title="Consolidation Notes",
        subtitle="How the three source decks were unified",
        bullets=[
            "App Architecture contributed role-based floor workflows and resilience design.",
            "Backend Architecture Flow contributed secure request lifecycle and layer governance.",
            "Web App Architecture contributed dashboards, lifecycle visibility, and planning surfaces.",
            "All three are preserved below as technical appendix snapshots for diligence review.",
        ],
    )
    add_contact_sheet_slide(
        prs,
        title="Technical Appendix Snapshot",
        subtitle="Factory Analytics App Architecture (13 source slides)",
        sheet_path=app_sheet,
    )
    add_contact_sheet_slide(
        prs,
        title="Technical Appendix Snapshot",
        subtitle="Factory Analytics Backend Architecture Flow (13 source slides)",
        sheet_path=backend_sheet,
    )
    add_contact_sheet_slide(
        prs,
        title="Technical Appendix Snapshot",
        subtitle="Factory Analytics Web App Architecture (13 source slides)",
        sheet_path=web_sheet,
    )
    add_bullet_slide(
        prs,
        title="Thank You",
        subtitle="Factory Analytics | Investor Discussion",
        bullets=[
            "Factory Analytics is built to convert shop-floor data into repeatable operational gains.",
            "The product architecture is deployment-ready, scalable, and role-anchored.",
            "Next step: align on pilot scope, target outcomes, and expansion milestones.",
        ],
    )

    prs.save(str(OUTPUT_PPTX))


if __name__ == "__main__":
    build_deck()
    print(f"Generated: {OUTPUT_PPTX}")
